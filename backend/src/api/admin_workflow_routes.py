"""Production admin workflow extensions: invitations, exports, strict policy import, and reassignment."""
from __future__ import annotations

import csv
import hashlib
import io
import secrets
from datetime import datetime, timedelta, timezone
from typing import Optional

from fastapi import APIRouter, Depends, File, HTTPException, Request, UploadFile
from fastapi.responses import StreamingResponse
from pydantic import BaseModel, Field
from sqlalchemy.orm import Session

from src.api.deps import require_role
from src.config import settings
from src.database.models import Adjuster, Claim, Policy, User
from src.database.hardening_models import AdjusterInvitation, ClaimAssignment, ClaimAuditEvent
from src.database.session import get_db
from src.utils.auth import get_password_hash
from src.utils.validators import validate_email, validate_full_name, validate_phone, CANONICAL_POLICY_TYPES

router = APIRouter(prefix="/api/v1/admin", tags=["Admin"])

class InviteAdjusterRequest(BaseModel):
    name: str = Field(..., min_length=2, max_length=100)
    email: str = Field(..., min_length=3, max_length=254)
    phone: str = Field(..., min_length=5, max_length=20)
    specialization: str = Field(...)

class StrictPolicyRequest(BaseModel):
    policy_number: str = Field(..., min_length=2, max_length=64)
    policy_type: str
    coverage_amount: float = Field(..., gt=0)
    deductible: float = Field(0, ge=0)
    effective_date: str
    expiry_date: str
    policyholder_name: str = Field(..., min_length=2, max_length=255)
    policyholder_dob: str
    policyholder_phone: str = Field(..., min_length=5, max_length=20)
    policyholder_email: Optional[str] = Field(None, max_length=254)

class ReassignClaimRequest(BaseModel):
    adjuster_id: str = Field(..., min_length=1)
    reason: str = Field(..., min_length=3, max_length=2000)

def _admin(request: Request, db: Session) -> User:
    return require_role(["ADMIN"])(request=request, current_user=__import__("src.api.deps", fromlist=["get_current_user"]).get_current_user(request=request, db=db))

def _parse_date(value: str, label: str):
    try:
        return datetime.strptime(value.strip(), "%Y-%m-%d").date()
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=f"{label} must be YYYY-MM-DD.") from exc

def _csv_or_xlsx(rows: list[list], headers: list[str], filename_base: str, fmt: str):
    if fmt.lower() == "xlsx":
        from openpyxl import Workbook
        bio = io.BytesIO()
        wb = Workbook()
        ws = wb.active
        ws.append(headers)
        for row in rows:
            ws.append(row)
        wb.save(bio)
        bio.seek(0)
        return StreamingResponse(bio, media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet", headers={"Content-Disposition": f'attachment; filename="{filename_base}.xlsx"'})
    out = io.StringIO()
    writer = csv.writer(out)
    writer.writerow(headers)
    writer.writerows(rows)
    return StreamingResponse(iter([out.getvalue()]), media_type="text/csv", headers={"Content-Disposition": f'attachment; filename="{filename_base}.csv"'})

@router.post("/adjusters/invite")
def invite_adjuster(payload: InviteAdjusterRequest, request: Request, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    try:
        name = validate_full_name(payload.name)
        email = validate_email(payload.email)
        phone = validate_phone(payload.phone)
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    specialization = payload.specialization.strip().lower()
    if specialization not in CANONICAL_POLICY_TYPES:
        raise HTTPException(status_code=400, detail=f"Specialization must be one of: {sorted(CANONICAL_POLICY_TYPES)}")
    if db.query(User).filter(User.email == email).first():
        raise HTTPException(status_code=409, detail="A user with this email already exists.")
    active_invite = db.query(AdjusterInvitation).filter(
        AdjusterInvitation.email == email,
        AdjusterInvitation.tenant_id == admin.tenant_id,
        AdjusterInvitation.accepted_at.is_(None),
        AdjusterInvitation.expires_at > datetime.now(timezone.utc),
    ).first()
    if active_invite:
        raise HTTPException(status_code=409, detail="An active invitation already exists for this email.")
    user_id = str(__import__("uuid").uuid4())
    user = User(
        id=user_id,
        tenant_id=str(admin.tenant_id or ""),
        full_name=name,
        email=email,
        phone=phone,
        password_hash=get_password_hash(secrets.token_urlsafe(48)),
        role="ADJUSTER",
        status="active",
        email_verified_at=None,
    )
    adjuster = Adjuster(
        id=user_id,
        user_id=user_id,
        tenant_id=str(admin.tenant_id or ""),
        name=name,
        email=email,
        phone=phone,
        specialization=specialization,
        claims_assigned=0,
        is_active=False,
    )
    raw_token = secrets.token_urlsafe(32)
    invite = AdjusterInvitation(
        tenant_id=str(admin.tenant_id or ""),
        email=email,
        name=name,
        phone=phone,
        specialization=specialization,
        token_hash=hashlib.sha256(raw_token.encode()).hexdigest(),
        expires_at=datetime.now(timezone.utc) + timedelta(days=3),
        created_by=admin.id,
    )
    db.add(user); db.add(adjuster); db.add(invite)
    try:
        db.commit(); db.refresh(invite)
    except Exception as exc:
        db.rollback()
        raise HTTPException(status_code=500, detail="Failed to create adjuster invitation.") from exc
    invite_url = f"{settings.PASSKEY_ORIGIN.rstrip('/')}/onboarding/adjuster?token={raw_token}"
    sent = False
    try:
        from src.utils.adjuster_invite_email import send_adjuster_invite_email
        sent = send_adjuster_invite_email(email, name, invite_url)
    except Exception:
        pass
    return {"id": invite.id, "adjuster_id": user_id, "status": "invited", "email": email, "invitation_url": invite_url, "email_sent": sent, "expires_at": invite.expires_at.isoformat()}

@router.get("/adjusters/invitations")
def list_invitations(request: Request, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    rows = db.query(AdjusterInvitation).filter(AdjusterInvitation.tenant_id == admin.tenant_id).order_by(AdjusterInvitation.created_at.desc()).all()
    now = datetime.now(timezone.utc)
    return {"items": [{"id": r.id, "email": r.email, "name": r.name, "specialization": r.specialization, "status": "accepted" if r.accepted_at else ("expired" if r.expires_at <= now else "invited"), "expires_at": r.expires_at.isoformat()} for r in rows]}

@router.get("/adjusters/export")
def export_adjusters(request: Request, format: str = "csv", is_active: Optional[bool] = None, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    q = db.query(Adjuster).filter(Adjuster.tenant_id == admin.tenant_id)
    if is_active is not None: q = q.filter(Adjuster.is_active == is_active)
    rows = q.order_by(Adjuster.name.asc()).all()
    headers=["name","email","phone","specialization","claims_assigned","is_active"]
    values=[[x.name,x.email,x.phone,x.specialization,x.claims_assigned,x.is_active] for x in rows]
    return _csv_or_xlsx(values, headers, "adjusters-export", format)

@router.get("/policies/export")
def export_policies(request: Request, format: str = "csv", policy_type: Optional[str] = None, is_active: Optional[bool] = None, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    q = db.query(Policy).filter(Policy.tenant_id == admin.tenant_id)
    if policy_type: q = q.filter(Policy.policy_type == policy_type.strip().lower())
    if is_active is not None: q = q.filter(Policy.is_active == is_active)
    rows = q.order_by(Policy.policy_number.asc()).all()
    headers=["policy_number","policy_type","coverage_amount","deductible","effective_date","expiry_date","is_active","policyholder_name","policyholder_dob","policyholder_phone","policyholder_email"]
    values=[[x.policy_number,x.policy_type,float(x.coverage_amount or 0),float(x.deductible or 0),str(x.effective_date),str(x.expiry_date),x.is_active,x.policyholder_name,str(x.policyholder_dob) if x.policyholder_dob else "",x.policyholder_phone,x.policyholder_email] for x in rows]
    return _csv_or_xlsx(values, headers, "policies-export", format)

@router.get("/policies/template")
def policy_template(request: Request, format: str = "csv", db: Session = Depends(get_db)):
    _admin(request, db)
    headers=["policy_number","policy_type","coverage_amount","deductible","effective_date","expiry_date","policyholder_name","policyholder_dob","policyholder_phone","policyholder_email"]
    sample=["POL-EXAMPLE-001","motor","1000000","10000","2026-01-01","2027-01-01","Example Policyholder","1990-01-01","9876543210","example@example.com"]
    return _csv_or_xlsx([sample], headers, "policy-import-template", format)

@router.post("/policies/strict")
def create_strict_policy(payload: StrictPolicyRequest, request: Request, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    ptype = payload.policy_type.strip().lower()
    if ptype not in CANONICAL_POLICY_TYPES:
        raise HTTPException(status_code=400, detail=f"Invalid policy_type '{payload.policy_type}'.")
    try:
        dob=_parse_date(payload.policyholder_dob,"Policyholder DOB")
        eff=_parse_date(payload.effective_date,"Effective date")
        exp=_parse_date(payload.expiry_date,"Expiry date")
        phone=validate_phone(payload.policyholder_phone)
        email=validate_email(payload.policyholder_email) if payload.policyholder_email else None
    except ValueError as exc:
        raise HTTPException(status_code=400, detail=str(exc)) from exc
    if exp <= eff: raise HTTPException(status_code=400, detail="Expiry date must be after effective date.")
    number=payload.policy_number.strip().upper()
    if db.query(Policy).filter(Policy.policy_number == number, Policy.tenant_id == admin.tenant_id).first():
        raise HTTPException(status_code=409, detail="Policy number already exists.")
    policy=Policy(id=str(__import__("uuid").uuid4()), tenant_id=str(admin.tenant_id or ""), policy_number=number, customer_id=None, policy_type=ptype, coverage_amount=payload.coverage_amount, deductible=payload.deductible, effective_date=eff, expiry_date=exp, is_active=True, policyholder_name=payload.policyholder_name.strip(), policyholder_dob=dob, policyholder_phone=phone, policyholder_email=email, policyholder_phone_last4=phone[-4:])
    db.add(policy); db.commit(); db.refresh(policy)
    return {"id":policy.id,"policy_number":policy.policy_number,"message":"Policy created successfully."}

@router.post("/policies/import-strict")
async def import_strict_policies(request: Request, file: UploadFile = File(...), db: Session = Depends(get_db)):
    admin = _admin(request, db)
    suffix=(file.filename or "").lower().rsplit(".",1)[-1]
    if suffix not in {"csv","xlsx"}: raise HTTPException(status_code=400, detail="Only .csv and .xlsx files are supported.")
    raw=await file.read(10*1024*1024+1)
    if len(raw)>10*1024*1024: raise HTTPException(status_code=413, detail="Policy import file exceeds 10 MB.")
    if suffix=="csv":
        reader=csv.DictReader(io.StringIO(raw.decode("utf-8-sig")))
    else:
        from openpyxl import load_workbook
        wb=load_workbook(io.BytesIO(raw),read_only=True,data_only=True); ws=wb.active; it=ws.iter_rows(values_only=True); heads=[str(x or "").strip() for x in next(it,None) or []]
        reader=({heads[i]:(vals[i] if i<len(vals) else "") for i in range(len(heads))} for vals in it)
    rows=list(reader)
    required=["policy_number","policy_type","coverage_amount","deductible","effective_date","expiry_date","policyholder_name","policyholder_dob","policyholder_phone"]
    fieldnames={str(x).strip().lower() for x in (reader.fieldnames or [])} if hasattr(reader,"fieldnames") else set(rows[0].keys() if rows else [])
    missing=[x for x in required if x not in fieldnames]
    if missing: raise HTTPException(status_code=400, detail={"missing_columns":missing})
    created=updated=0; errors=[]
    for idx, raw_row in enumerate(rows, start=2):
        row={str(k).strip().lower(): ("" if v is None else str(v).strip()) for k,v in raw_row.items() if k is not None}
        try:
            number=row.get("policy_number","").upper()
            ptype=row.get("policy_type","").lower()
            if not number: raise ValueError("Missing policy_number")
            if ptype not in CANONICAL_POLICY_TYPES: raise ValueError(f"Invalid policy_type '{ptype}'")
            cov=float(row.get("coverage_amount","")); ded=float(row.get("deductible","0"))
            eff=_parse_date(row.get("effective_date",""),"Effective date"); exp=_parse_date(row.get("expiry_date",""),"Expiry date")
            name=validate_full_name(row.get("policyholder_name","")); dob=_parse_date(row.get("policyholder_dob",""),"Policyholder DOB")
            phone=validate_phone(row.get("policyholder_phone",""))
            email=validate_email(row.get("policyholder_email")) if row.get("policyholder_email") else None
            if exp <= eff: raise ValueError("Expiry date must be after effective date")
            existing=db.query(Policy).filter(Policy.policy_number==number,Policy.tenant_id==admin.tenant_id).first()
            if existing:
                existing.policy_type=ptype; existing.coverage_amount=cov; existing.deductible=ded; existing.effective_date=eff; existing.expiry_date=exp; existing.policyholder_name=name; existing.policyholder_dob=dob; existing.policyholder_phone=phone; existing.policyholder_email=email; existing.policyholder_phone_last4=phone[-4:]; updated+=1
            else:
                db.add(Policy(id=str(__import__("uuid").uuid4()),tenant_id=str(admin.tenant_id or ""),policy_number=number,customer_id=None,policy_type=ptype,coverage_amount=cov,deductible=ded,effective_date=eff,expiry_date=exp,is_active=True,policyholder_name=name,policyholder_dob=dob,policyholder_phone=phone,policyholder_email=email,policyholder_phone_last4=phone[-4:])); created+=1
        except (ValueError, TypeError) as exc:
            errors.append({"row":idx,"policy_number":row.get("policy_number"),"error":str(exc)})
    db.commit()
    return {"created":created,"updated":updated,"total_processed":created+updated,"errors":errors}

@router.post("/claims/{ticket_id}/reassign")
def reassign_claim(ticket_id: str, payload: ReassignClaimRequest, request: Request, db: Session = Depends(get_db)):
    admin = _admin(request, db)
    claim=db.query(Claim).filter(Claim.ticket_id==ticket_id,Claim.tenant_id==admin.tenant_id).first()
    if not claim: raise HTTPException(status_code=404,detail="Claim not found.")
    target=db.query(Adjuster).filter(Adjuster.id==payload.adjuster_id,Adjuster.tenant_id==admin.tenant_id,Adjuster.is_active.is_(True)).first()
    if not target: raise HTTPException(status_code=404,detail="Active target adjuster not found.")
    active=db.query(ClaimAssignment).filter(ClaimAssignment.claim_id==claim.id,ClaimAssignment.tenant_id==claim.tenant_id,ClaimAssignment.is_active.is_(True)).first()
    if active and str(active.adjuster_id)==str(target.id): return {"success":True,"already_assigned":True,"adjuster_id":target.id}
    if active:
        old=db.query(Adjuster).filter(Adjuster.id==active.adjuster_id,Adjuster.tenant_id==claim.tenant_id).first()
        active.is_active=False; active.unassigned_at=datetime.now(timezone.utc)
        if old: old.claims_assigned=max(0,int(old.claims_assigned or 0)-1)
    if str(target.id)!=(str(active.adjuster_id) if active else ""):
        target.claims_assigned=int(target.claims_assigned or 0)+1
        db.add(ClaimAssignment(claim_id=claim.id,tenant_id=str(claim.tenant_id or ""),adjuster_id=target.id,assigned_by=admin.id,reason=payload.reason,is_active=True))
    state=dict(claim.pipeline_state or {}); state["assigned_adjuster_id"]=target.id; state["assigned_adjuster_name"]=target.name; claim.pipeline_state=state
    db.add(ClaimAuditEvent(claim_id=claim.id,tenant_id=str(claim.tenant_id or ""),actor_user_id=admin.id,event_type="claim_reassigned",new_value_json={"adjuster_id":target.id,"adjuster_name":target.name},reason=payload.reason))
    db.commit()
    return {"success":True,"already_assigned":False,"adjuster_id":target.id,"adjuster_name":target.name}

